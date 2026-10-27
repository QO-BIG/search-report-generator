"""Build the Search Report .docx from structured data.

Opens assets/template.docx (QO letterhead header, footer and A4 section, body
emptied) and writes the firm's Search Report template into it:
  - details typed in the app (ref no., date, client, class, goods...) are filled in;
  - the prior-marks table is filled from the Prior Marks PDF;
  - the judgement parts (summary table, Part I analysis, Opinion, similarity
    column) are left empty for staff to type in Word.
Everything stays editable Word content.
"""
from __future__ import annotations

import base64
import copy
import io
import os
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from PIL import Image

ASSETS = Path(__file__).resolve().parent / "assets"
TEMPLATE_PATH = ASSETS / "template.docx"
SIGNATURE_PATH = ASSETS / "signature.png"  # local only; kept out of the public repo
FONT = "Times New Roman"
RED = RGBColor(0xFF, 0x00, 0x00)

# numbering definitions present in the template
NUM_ROMAN = 1  # I) II)
NUM_TABLE = 3  # 1. 2. 3. in the "No." column

# ------------------------------------------------------------------ fixed text
RE_PREFIX = "TRADEMARK SEARCH ON MARK “"
INTRO = ("We refer to your instruction to conduct a trademark search on the above mark and wish "
         "to thank you for entrusting us with this assignment.")
SUMMARY_LABEL = "Summary of search:"
SUMMARY_HEADERS = ["Availability based on consideration of prior marks",
                   "Inherent Registrability (distinctive/non-descriptive)", "Chances of Registration"]
FOOTNOTE = ("*Please note that whether an application is accepted/ rejected is ultimately subject to "
            "the discretion of the examiner, and there will be additional fees to respond.")
ASPECTS = "Kindly note that for the subject mark, the search report covers the following two (2) aspects:"
ASPECT_ITEMS = [("Absolute grounds:", "Whether the mark is inherently registrable; and"),
                ("Relative grounds:", "Whether there are any identical or similar marks.")]
PART1_HEADING = "PART I – Absolute grounds"
PART1_INTRO = ("One of the significant factors in determining the registrability of a trademark is the "
               "question of whether the mark is sufficiently distinctive. If the Registrar finds that an "
               "applied-for mark lacks sufficient distinctiveness, the Registrar will reject the "
               "application on absolute grounds.")
PART1_CLOSE = ("Nevertheless, the decision to accept or reject a mark ultimately rests with the "
               "Registrar's discretion. Please note that there may be additional fees incurred in "
               "responding to and overcoming any objections.")
PART2_HEADING = "PART II – Relative grounds"
PART2_INTRO = ("Trademark applications are rejected on ‘relative grounds’ when the trademark conflicts "
               "with prior trademark rights. Having two identical (or very similar) trademarks for the "
               "same or related goods and/or services could cause confusion among consumers. If the "
               "trademark is considered to be identical or confusingly similar to an existing prior mark, "
               "for having the same and/or related goods and/or services, it will be rejected by the "
               "Registrar.")
SEARCH_PARA = ("We have completed the search on the above subject mark in {classes} and its related classes "
               "in the Trademark Register of Malaysia. The search covers registered trademark applications "
               "available for public search as of {date}. The search result is summarised below.")
INTENDED_LABEL = "The intended {kind} for application by the Applicant are as follows:"
MARK_HEADING_SUFFIX = " & RELATED CLASS(ES)"
LEAD_IN = "We found these prior marks that are similar to the subject mark."
TABLE_HEADERS = ["No.", "Mark", "Application No.", "Class", "Status", "Similarity of Goods / Services"]
TABLE_GRID = [680, 2468, 1843, 992, 1843, 1203]  # = the text width (9029), so both page margins match
ATTACH = "We attach herewith a copy of the printout of the prior marks above for your retention."
OPINION_CLOSE = ("Notwithstanding the above opinion, the decision to accept or to reject the application "
                 "ultimately lies within the discretion of the Registrar. Please be advised that in the "
                 "event of an objection or office action, there will be additional fees incurred to "
                 "respond to the objection.")
SCOPE_HEADING = "Scope of Search"
SCOPE_TEXT = ("Please note that the above search is conducted on the database provided based on the latest "
              "records made available to the public at the Registry of Trademark, which may be lacking a "
              "complete list of live trademarks in Malaysia. Due to the condition of the records in the "
              "database, the search cannot be taken as conclusive. Please do allow room for human errors "
              "should any identical or confusingly similar marks have been inadvertently missed out during "
              "the search.")
PART1_REMINDER = "[Write the absolute grounds analysis here]"
OPINION_REMINDER = "[Write the opinion for Class {cls} here]"
THANKS = "We thank you for your support and look forward to receiving further instructions on the matter."


# ------------------------------------------------------------------ data model
@dataclass
class TableRow:
    """One row of the prior-marks table. Several numbers = merged row."""
    numbers: list[str]
    classes: list[str]  # one entry per number, e.g. ["6"], or ["9, 35, 38"]
    status: str
    image_png: bytes | None = None
    mark_text: str = ""


@dataclass
class ClassSection:
    """The prior-marks table and Opinion for one searched class."""
    cls: int
    rows: list[TableRow]


@dataclass
class ReportData:
    ref_no: str
    letter_date: date
    attention: str
    company: str
    address_lines: list[str]
    salutation: str
    subject_text: str
    classes: list[int]
    intended_goods: dict[int, str]
    search_date: date
    sections: list[ClassSection]
    subject_image: bytes | None = None
    signatory: str = "TAN SIN SU"
    signature_png: bytes | None = None
    company_name: str = "Quality Oracle Sdn Bhd"


# ------------------------------------------------------------------ helpers
def ordinal(n: int) -> str:
    suffix = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def class_no(c: int) -> str:
    """Two-digit class number as in the template: 5 -> '05'."""
    return f"{c:02d}"


def classes_phrase(classes: list[int], word: str = "Class") -> str:
    """'Class 05' / 'Classes 05 & 10' / 'Classes 05, 10 & 44'."""
    nums = [class_no(c) for c in classes]
    if len(nums) <= 1:
        return f"{word} {''.join(nums)}"
    return f"{word}es {', '.join(nums[:-1])} & {nums[-1]}"


def goods_kind(classes: list[int]) -> str:
    """Classes 1-34 are goods, 35-45 services."""
    kinds = {"goods" if c <= 34 else "services" for c in classes}
    return kinds.pop() if len(kinds) == 1 else "goods/services"


def default_signature() -> bytes | None:
    """Signature image: SIGNATURE_PNG_BASE64 (env / app secrets) or the local assets file."""
    b64 = os.environ.get("SIGNATURE_PNG_BASE64")
    if b64:
        return base64.b64decode(b64)
    return SIGNATURE_PATH.read_bytes() if SIGNATURE_PATH.exists() else None


def _fit(png: bytes, max_w_in: float, max_h_in: float) -> tuple[Emu, Emu]:
    with Image.open(io.BytesIO(png)) as im:
        w, h = im.size
    scale = min(max_w_in / w, max_h_in / h)
    return Inches(w * scale), Inches(h * scale)


def _set_fonts(run, bold=None, italic=None, underline=None, superscript=False, size=None):
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if superscript:
        run.font.superscript = True
    if size:
        run.font.size = Pt(size)


# CT_PPr children that must follow w:pBdr (schema order)
_PPR_AFTER_PBDR = tuple("w:" + t for t in (
    "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct",
    "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"))


class Builder:
    def __init__(self, data: ReportData):
        self.d = data
        self.doc = Document(str(TEMPLATE_PATH))
        self.body = self.doc.element.body

    # -- low level ------------------------------------------------------
    def _new_p(self, *, jc="both", line=None, before=None, after=None,
               ind_left=None, ind_hanging=None, keep_next=False):
        p = self.doc.add_paragraph()  # python-docx inserts before sectPr
        pf = p.paragraph_format
        if jc == "both":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif jc == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if line:
            pf.line_spacing = line
        if before is not None:
            pf.space_before = Pt(before)
        if after is not None:
            pf.space_after = Pt(after)
        if ind_left is not None:
            pf.left_indent = Emu(ind_left * 635)
        if ind_hanging is not None:
            pf.first_line_indent = Emu(-ind_hanging * 635)
        if keep_next:
            pf.keep_with_next = True
        self._mark_font(p)
        return p

    @staticmethod
    def _mark_font(p):
        """Paragraph-mark font (affects list numbers and empty-line height)."""
        ppr = p._p.get_or_add_pPr()
        rpr = ppr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ppr.insert_element_before(rpr, "w:sectPr", "w:pPrChange")
        fonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), FONT)
        rpr.insert(0, fonts)

    def _run(self, p, text, **fmt):
        r = p.add_run(text)
        _set_fonts(r, **fmt)
        return r

    def _picture(self, p, png: bytes, max_w: float, max_h: float, **fmt):
        w, h = _fit(png, max_w, max_h)
        r = p.add_run()
        _set_fonts(r, **fmt)
        r.add_picture(io.BytesIO(png), width=w, height=h)
        return r

    def blank(self, line=1.15):
        return self._new_p(line=line)

    def text(self, s: str, **fmt):
        p = self._new_p(line=1.15)
        self._run(p, s, **fmt)
        return p

    def staff_space(self, reminder: str):
        """Red reminder line where staff type their own analysis in Word (red in the template)."""
        self.blank()
        p = self._new_p(line=1.15)
        r = self._run(p, reminder)
        r.font.color.rgb = RED
        self.blank()

    def _numbered(self, p, num_id: int):
        numpr = p._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = num_id

    def _bottom_border(self, p):
        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "auto")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        ppr.insert_element_before(bdr, *_PPR_AFTER_PBDR)

    def _subject(self, p, **fmt):
        """The subject mark: its image if one was given, else its name."""
        if self.d.subject_image:
            self._picture(p, self.d.subject_image, 1.6, 0.66, **fmt)
        else:
            self._run(p, self.d.subject_text.upper(), **fmt)

    def _date_runs(self, p, dt: date, **fmt):
        o = ordinal(dt.day)
        self._run(p, o[:-2], **fmt)
        self._run(p, o[-2:], superscript=True, **fmt)
        self._run(p, f" {dt.strftime('%B %Y')}", **fmt)

    # -- tables ---------------------------------------------------------
    def _table(self, grid: list[int], *, width=None, ind=None, fixed=False, center=False):
        tbl = self.doc.add_table(rows=0, cols=len(grid))
        tblpr = tbl._tbl.tblPr
        for el in list(tblpr):
            if el.tag in (qn("w:tblW"), qn("w:tblStyle"), qn("w:tblLook")):
                tblpr.remove(el)
        tblw = OxmlElement("w:tblW")
        tblw.set(qn("w:w"), str(width or 0))
        tblw.set(qn("w:type"), "dxa" if width else "auto")
        tblpr.append(tblw)
        if center:
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), "center")
            tblpr.append(jc)
        if ind is not None:
            ti = OxmlElement("w:tblInd")
            ti.set(qn("w:w"), str(ind))
            ti.set(qn("w:type"), "dxa")
            tblpr.append(ti)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "auto")):
                b.set(qn(k), v)
            borders.append(b)
        tblpr.append(borders)
        if fixed:
            lay = OxmlElement("w:tblLayout")
            lay.set(qn("w:type"), "fixed")
            tblpr.append(lay)
        look = OxmlElement("w:tblLook")
        for k, v in (("w:val", "04A0"), ("w:firstRow", "1"), ("w:lastRow", "0"), ("w:firstColumn", "1"),
                     ("w:lastColumn", "0"), ("w:noHBand", "0"), ("w:noVBand", "1")):
            look.set(qn(k), v)
        tblpr.append(look)
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), grid):
            gc.set(qn("w:w"), str(w))
        return tbl

    def _row(self, tbl, grid, *, header=False, height=None):
        row = tbl.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if header:
            trpr.append(OxmlElement("w:tblHeader"))
        if height:
            th = OxmlElement("w:trHeight")
            th.set(qn("w:val"), str(height))
            trpr.append(th)
        for cell, w in zip(row.cells, grid):
            cell.width = Emu(w * 635)
        return row

    @staticmethod
    def _vmerge(cell, restart: bool):
        vm = cell._tc.get_or_add_tcPr().get_or_add_vMerge()
        if restart:
            vm.set(qn("w:val"), "restart")

    def _cell_p(self, cell, *, jc="center", before=None, after=None):
        p = cell.paragraphs[0]
        self._mark_font(p)
        if jc == "both":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif jc == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    # -- sections ---------------------------------------------------------
    def letter_head(self):
        d = self.d
        for _ in range(2):
            self._new_p(line=2.0)
        p = self._new_p(line=2.0)
        self._run(p, "By E-mail Only", bold=True, italic=True)
        p = self._new_p()
        self._run(p, "Our Ref. No.:", italic=True)
        self._run(p, f" {d.ref_no}")
        self._new_p()
        p = self._new_p()
        self._date_runs(p, d.letter_date)
        self._new_p()
        p = self._new_p(jc=None, line=1.15)
        self._run(p, f"Attention: {d.attention}", bold=True)
        if d.company:
            p = self._new_p(jc=None, line=1.15)
            self._run(p, d.company, bold=True)
        for line in d.address_lines:
            if line.strip():
                p = self._new_p(jc=None, line=1.15)
                self._run(p, line.strip())
        self._new_p()
        p = self._new_p()
        self._run(p, d.salutation)
        self._new_p()

    def re_line(self):
        p = self._new_p(ind_left=720, ind_hanging=720)
        self._bottom_border(p)
        self._run(p, "RE:", bold=True)
        self._run(p, "\t" + RE_PREFIX, bold=True)
        self._subject(p, bold=True)
        self._run(p, f"” IN {classes_phrase(self.d.classes).upper()} IN MALAYSIA", bold=True)
        self._new_p(ind_left=720, ind_hanging=720)

    def summary(self):
        self.text(INTRO)
        p = self._new_p(line=1.15, before=6)
        self._run(p, SUMMARY_LABEL)
        # one class: as the template; several classes: a Class column and one row per class (as past reports)
        multi = len(self.d.classes) > 1
        grid = [1129, 2971, 2774, 2145] if multi else [2977, 2835, 2191]
        headers = (["Class"] if multi else []) + SUMMARY_HEADERS
        tbl = self._table(grid, center=True, ind=0)
        row = self._row(tbl, grid, header=True)
        for cell, h in zip(row.cells, headers):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = self._cell_p(cell, before=3, after=3)
            self._run(p, h, bold=True)
        for c in self.d.classes if multi else [None]:
            row = self._row(tbl, grid, height=500)  # assessment cells are left empty for staff
            for i, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = self._cell_p(cell, before=3, after=3)
                if multi and i == 0:
                    self._run(p, class_no(c))
        self.text(FOOTNOTE, bold=True, italic=True)
        self._new_p()
        self.text(ASPECTS)
        for label, text in ASPECT_ITEMS:
            p = self._new_p(line=1.15, ind_left=851, ind_hanging=470)
            self._numbered(p, NUM_ROMAN)
            self._run(p, label, bold=True)
            self._run(p, "\t" + text)
        self._new_p()

    def part1(self):
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, PART1_HEADING, bold=True, underline=True)
        p = self._new_p(line=1.15, before=6)
        self._run(p, PART1_INTRO)
        self.staff_space(PART1_REMINDER)
        self.text(PART1_CLOSE)
        self.blank()

    def part2(self):
        d = self.d
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, PART2_HEADING, bold=True, underline=True)
        p = self._new_p(line=1.15, before=6)
        self._run(p, PART2_INTRO)
        self.blank()

        before, after = SEARCH_PARA.split("{date}")
        p = self._new_p(line=1.15)
        self._run(p, before.format(classes=classes_phrase(d.classes, "class")))
        self._date_runs(p, d.search_date)
        self._run(p, after)
        self.blank()

        self.text(INTENDED_LABEL.format(kind=goods_kind(d.classes)))
        for c in d.classes:
            p = self._new_p(line=1.15, before=6)
            self._run(p, f"Class {class_no(c)}", bold=True, italic=True, underline=True)
            self.text(d.intended_goods.get(c, "").strip())
        self.blank()

        for sec in d.sections:
            self.class_section(sec)

    def class_section(self, sec: ClassSection):
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, "MARK “", bold=True, underline=True)
        self._subject(p, bold=True, underline=True)
        self._run(p, f"” IN {classes_phrase([sec.cls]).upper()}{MARK_HEADING_SUFFIX}", bold=True, underline=True)
        p = self._new_p(line=1.15, before=6, after=6, keep_next=True)
        self._run(p, LEAD_IN)
        self.marks_table(sec.rows)
        self.blank()
        self.text(ATTACH)
        self.blank()
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, "Opinion", bold=True, underline=True)
        self.staff_space(OPINION_REMINDER.format(cls=class_no(sec.cls)))
        self.text(OPINION_CLOSE)
        self.blank()

    def _restarted_num(self, num_id: int) -> int:
        """A separate copy of num_id's list, so each table counts from 1.

        (A startOverride on a shared list makes Word's PDF export restart the count
        at the page break, so each table gets its own list definition instead.)
        """
        numbering = self.doc.part.numbering_part.element
        base = next(n for n in numbering.num_lst if n.numId == num_id)
        abstracts = numbering.findall(qn("w:abstractNum"))
        src = next(a for a in abstracts if a.get(qn("w:abstractNumId")) == str(base.abstractNumId.val))
        new_id = max(int(a.get(qn("w:abstractNumId"))) for a in abstracts) + 1
        abstract = copy.deepcopy(src)
        abstract.set(qn("w:abstractNumId"), str(new_id))
        nsid = abstract.find(qn("w:nsid"))
        if nsid is not None:
            nsid.set(qn("w:val"), f"{0x5E000000 + new_id:08X}")
        abstracts[-1].addnext(abstract)  # abstractNums must precede the nums
        return numbering.add_num(new_id).numId

    def marks_table(self, rows: list[TableRow]):
        grid = TABLE_GRID
        num_id = self._restarted_num(NUM_TABLE)
        tbl = self._table(grid, width=sum(grid), fixed=True, center=True)
        row = self._row(tbl, grid, header=True)
        for i, (cell, h) in enumerate(zip(row.cells, TABLE_HEADERS)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = self._cell_p(cell, jc="center" if i in (4, 5) else "both")
            # the long "Similarity of Goods / Services" header is 11 pt in the template
            self._run(p, h, bold=True, size=11 if len(h) > 20 else None)

        for tr in rows or [TableRow(numbers=[], classes=[], status="")]:
            n = max(1, len(tr.numbers))
            for k in range(n):
                height = (624 if tr.image_png else 397) if n == 1 else 390
                row = self._row(tbl, grid, height=height)
                cells = row.cells
                for c in cells:
                    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if n > 1:
                    for ci in (0, 1, 4, 5):
                        self._vmerge(cells[ci], restart=(k == 0))
                if k == 0:
                    p = self._cell_p(cells[0], jc="both")
                    self._numbered(p, num_id)
                    p = self._cell_p(cells[1], before=2, after=2)
                    if tr.image_png:
                        self._picture(p, tr.image_png, 1.22, 0.72)  # as staff size them: wide ~1.2", logos ~0.7"
                    else:
                        self._run(p, tr.mark_text)
                    p = self._cell_p(cells[4])
                    self._run(p, tr.status)
                    self._cell_p(cells[5])  # similarity: typed by staff
                p = self._cell_p(cells[2])
                self._run(p, tr.numbers[k] if tr.numbers else "")
                p = self._cell_p(cells[3])
                p.paragraph_format.right_indent = Emu(-112 * 635)
                self._run(p, tr.classes[k] if k < len(tr.classes) else "")
                if k < n - 1:  # keep a merged group on one page
                    for c in cells:
                        for cp in c.paragraphs:
                            cp.paragraph_format.keep_with_next = True

    def closing(self):
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, SCOPE_HEADING, bold=True, underline=True)
        p = self._new_p(line=1.15, before=6)
        self._run(p, SCOPE_TEXT)
        self.blank()
        self.blank()
        p = self._new_p(line=1.15, keep_next=True)  # keep the thanks line with the signature block
        self._run(p, THANKS)
        self._new_p(line=1.15, keep_next=True)
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, "Yours sincerely,")
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, self.d.company_name, bold=True)
        p = self._new_p(line=1.15, keep_next=True)
        sig = self.d.signature_png if self.d.signature_png is not None else default_signature()
        if sig:
            r = self._picture(p, sig, 1.05, 0.55)
            self._inline_to_anchor(r)
        self._new_p(line=1.15, keep_next=True)
        p = self._new_p(line=1.15, keep_next=True)
        self._run(p, "……………………")
        p = self._new_p(line=1.15)
        self._run(p, self.d.signatory, bold=True)
        p = self._new_p(line=1.15)
        self._run(p, "Encls.")

    def _inline_to_anchor(self, run):
        """Turn the signature into a floating 'behind text' picture like the template."""
        inline = run._r.find(".//" + qn("wp:inline"))
        if inline is None:
            return
        anchor = OxmlElement("wp:anchor")
        for k, v in (("distT", "0"), ("distB", "0"), ("distL", "114300"), ("distR", "114300"),
                     ("simplePos", "0"), ("relativeHeight", "251658240"), ("behindDoc", "1"),
                     ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1")):
            anchor.set(k, v)
        sp = OxmlElement("wp:simplePos")
        sp.set("x", "0")
        sp.set("y", "0")
        anchor.append(sp)
        for tag, rel in (("wp:positionH", "column"), ("wp:positionV", "paragraph")):
            pos = OxmlElement(tag)
            pos.set("relativeFrom", rel)
            off = OxmlElement("wp:posOffset")
            off.text = "0"
            pos.append(off)
            anchor.append(pos)
        anchor.append(copy.deepcopy(inline.find(qn("wp:extent"))))
        eff = OxmlElement("wp:effectExtent")
        for k in ("l", "t", "r", "b"):
            eff.set(k, "0")
        anchor.append(eff)
        anchor.append(OxmlElement("wp:wrapNone"))
        for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
            el = inline.find(qn(tag))
            if el is not None:
                anchor.append(copy.deepcopy(el))
        inline.getparent().replace(inline, anchor)

    def _unique_docpr_ids(self):
        """python-docx numbers pictures from 1 per run; make docPr ids unique."""
        for i, el in enumerate(self.body.iter(qn("wp:docPr")), start=1):
            el.set("id", str(i))

    def build(self) -> bytes:
        self.letter_head()
        self.re_line()
        self.summary()
        self.part1()
        self.part2()
        self.closing()
        self._unique_docpr_ids()
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def build_report(data: ReportData) -> bytes:
    return Builder(data).build()


def report_filename(subject_text: str) -> str:
    safe = re.sub(r'[\\/:*?"<>|]+', "", subject_text).strip() or "Mark"
    return f"[MY] Search Report - {safe}.docx"
